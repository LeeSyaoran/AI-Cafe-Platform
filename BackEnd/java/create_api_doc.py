from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ========== TITLE ==========
title = doc.add_heading('AI CAFE PLATFORM - API DOCUMENTATION', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph('Postman Collection Guide')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Base URL info
doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('Base URL: ').bold = True
p.add_run('https://api-staging.aicafe.vn/api')

# Environment info
doc.add_paragraph()
p = doc.add_paragraph()
p.add_run('Production: ').bold = True
p.add_run('https://api.aicafe.vn/api')
p = doc.add_paragraph()
p.add_run('Staging: ').bold = True
p.add_run('https://api-staging.aicafe.vn/api')
p = doc.add_paragraph()
p.add_run('Local: ').bold = True
p.add_run('http://localhost:3000/api')

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 1. AUTHENTICATION
# ============================================================
doc.add_heading('1. AUTHENTICATION', level=1)

# --- POST /v1/auth/send-otp ---
doc.add_heading('POST /v1/auth/send-otp', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Send OTP to phone or email')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code1 = '''{
    "email": "test@example.com",
    "phone": "0909123456",
    "type": "register"
}'''
p = doc.add_paragraph(code1)
p.style = 'Quote'

doc.add_paragraph()

# --- POST /v1/auth/verify-otp ---
doc.add_heading('POST /v1/auth/verify-otp', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Verify OTP and get tokens')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code2 = '''{
    "email": "test@example.com",
    "phone": "0909123456",
    "otp": "123456"
}'''
p = doc.add_paragraph(code2)
p.style = 'Quote'

doc.add_paragraph()

# --- POST /v1/auth/register ---
doc.add_heading('POST /v1/auth/register', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Register new user account')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code3 = '''{
    "email": "test@example.com",
    "password": "password123",
    "phone": "0909123456"
}'''
p = doc.add_paragraph(code3)
p.style = 'Quote'

doc.add_paragraph()

# --- POST /v1/auth/login ---
doc.add_heading('POST /v1/auth/login', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Login with email and password')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code4 = '''{
    "email": "test@example.com",
    "password": "password123"
}'''
p = doc.add_paragraph(code4)
p.style = 'Quote'

doc.add_paragraph()

# --- POST /v1/auth/refresh ---
doc.add_heading('POST /v1/auth/refresh', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Refresh access token')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code5 = '''{
    "refreshToken": "{{refreshToken}}"
}'''
p = doc.add_paragraph(code5)
p.style = 'Quote'

doc.add_paragraph()

# --- POST /v1/auth/logout ---
doc.add_heading('POST /v1/auth/logout', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Logout and invalidate token')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code6 = '''{
    "refreshToken": "{{refreshToken}}"
}'''
p = doc.add_paragraph(code6)
p.style = 'Quote'

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 2. USER PROFILE
# ============================================================
doc.add_heading('2. USER PROFILE', level=1)

# --- GET /v1/me ---
doc.add_heading('GET /v1/me', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get current user profile')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- PUT /v1/me ---
doc.add_heading('PUT /v1/me', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Update current user profile')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code7 = '''{
    "fullName": "Nguyen Van Test",
    "displayName": "TestUser",
    "dateOfBirth": "1990-01-15",
    "gender": "male"
}'''
p = doc.add_paragraph(code7)
p.style = 'Quote'

doc.add_paragraph()

# --- GET /v1/users/{id} ---
doc.add_heading('GET /v1/users/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get user by ID')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 3. CART
# ============================================================
doc.add_heading('3. CART', level=1)

# --- GET /v1/cart ---
doc.add_heading('GET /v1/cart', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get user cart for a cafe')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('cafeId: UUID (required)')

doc.add_paragraph()

# --- POST /v1/cart/items ---
doc.add_heading('POST /v1/cart/items', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Add item to cart')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code8 = '''{
    "cafeId": "REPLACE_WITH_CAFE_ID",
    "productId": "REPLACE_WITH_PRODUCT_ID",
    "variantId": null,
    "quantity": 1,
    "options": {},
    "modifiers": [],
    "notes": ""
}'''
p = doc.add_paragraph(code8)
p.style = 'Quote'

doc.add_paragraph()

# --- PUT /v1/cart/items/{id} ---
doc.add_heading('PUT /v1/cart/items/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Update cart item')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code9 = '''{
    "quantity": 2,
    "notes": "Extra ice"
}'''
p = doc.add_paragraph(code9)
p.style = 'Quote'

doc.add_paragraph()

# --- DELETE /v1/cart/items/{id} ---
doc.add_heading('DELETE /v1/cart/items/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Remove item from cart')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- DELETE /v1/cart ---
doc.add_heading('DELETE /v1/cart', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Clear all items from cart')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('cafeId: UUID (required)')

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 4. ORDERS
# ============================================================
doc.add_heading('4. ORDERS', level=1)

# --- POST /v1/orders ---
doc.add_heading('POST /v1/orders', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Create new order from cart')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code10 = '''{
    "companyId": "REPLACE_WITH_COMPANY_ID",
    "cafeId": "REPLACE_WITH_CAFE_ID",
    "orderType": "pickup",
    "deliveryAddressId": null,
    "customerNote": "",
    "promoCode": ""
}'''
p = doc.add_paragraph(code10)
p.style = 'Quote'

doc.add_paragraph()

# --- GET /v1/orders ---
doc.add_heading('GET /v1/orders', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get user order history')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('page: int (default: 0)')
doc.add_paragraph('size: int (default: 20)')

doc.add_paragraph()

# --- GET /v1/orders/{id} ---
doc.add_heading('GET /v1/orders/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get order by ID')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()

# --- GET /v1/orders/number/{orderNumber} ---
doc.add_heading('GET /v1/orders/number/{orderNumber}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get order by order number')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()

# --- PUT /v1/orders/{id}/confirm ---
doc.add_heading('PUT /v1/orders/{id}/confirm', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Staff confirms order')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()

# --- PUT /v1/orders/{id}/preparing ---
doc.add_heading('PUT /v1/orders/{id}/preparing', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Staff starts preparing order')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()

# --- PUT /v1/orders/{id}/ready ---
doc.add_heading('PUT /v1/orders/{id}/ready', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Staff marks order as ready')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()

# --- PUT /v1/orders/{id}/complete ---
doc.add_heading('PUT /v1/orders/{id}/complete', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Staff completes order')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()

# --- POST /v1/orders/{id}/cancel ---
doc.add_heading('POST /v1/orders/{id}/cancel', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Cancel order')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code11 = '''{
    "reason": "Customer requested cancellation"
}'''
p = doc.add_paragraph(code11)
p.style = 'Quote'

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 5. PAYMENTS
# ============================================================
doc.add_heading('5. PAYMENTS', level=1)

# --- GET /v1/payments/methods ---
doc.add_heading('GET /v1/payments/methods', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get available payment methods')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- POST /v1/payments ---
doc.add_heading('POST /v1/payments', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Create payment for order')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code12 = '''{
    "orderId": "REPLACE_WITH_ORDER_ID",
    "companyId": "REPLACE_WITH_COMPANY_ID",
    "amount": 50000,
    "paymentMethod": "CREDIT",
    "returnUrl": "https://aicafe.vn/payment/return"
}'''
p = doc.add_paragraph(code12)
p.style = 'Quote'

doc.add_paragraph()

# --- GET /v1/payments/{id} ---
doc.add_heading('GET /v1/payments/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get payment by ID')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- GET /v1/payments/order/{orderId} ---
doc.add_heading('GET /v1/payments/order/{orderId}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get payment by order ID')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- POST /v1/payments/{id}/cancel ---
doc.add_heading('POST /v1/payments/{id}/cancel', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Cancel pending payment')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 6. LOYALTY & REWARDS
# ============================================================
doc.add_heading('6. LOYALTY & REWARDS', level=1)

# --- GET /v1/loyalty/account ---
doc.add_heading('GET /v1/loyalty/account', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get user loyalty account info')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- GET /v1/loyalty/rewards ---
doc.add_heading('GET /v1/loyalty/rewards', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get available rewards for redemption')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- POST /v1/loyalty/redeem ---
doc.add_heading('POST /v1/loyalty/redeem', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Redeem loyalty points for reward')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code13 = '''{
    "rewardId": "REPLACE_WITH_REWARD_ID",
    "orderId": "REPLACE_WITH_ORDER_ID"
}'''
p = doc.add_paragraph(code13)
p.style = 'Quote'

doc.add_paragraph()

# --- GET /v1/loyalty/history ---
doc.add_heading('GET /v1/loyalty/history', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get loyalty transaction history')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('limit: int (default: 50)')

doc.add_paragraph()

# --- GET /v1/loyalty/redemptions ---
doc.add_heading('GET /v1/loyalty/redemptions', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get reward redemption history')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 7. SUBSCRIPTIONS
# ============================================================
doc.add_heading('7. SUBSCRIPTIONS', level=1)

# --- GET /v1/subscriptions/plans ---
doc.add_heading('GET /v1/subscriptions/plans', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get available membership plans')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()

# --- GET /v1/subscriptions/plans/{id} ---
doc.add_heading('GET /v1/subscriptions/plans/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get subscription plan details')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()

# --- GET /v1/subscriptions/me ---
doc.add_heading('GET /v1/subscriptions/me', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get user current subscription')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- POST /v1/subscriptions ---
doc.add_heading('POST /v1/subscriptions', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Subscribe to membership plan')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code14 = '''{
    "planId": "REPLACE_WITH_PLAN_ID",
    "billingCycle": "monthly",
    "paymentMethod": "CREDIT"
}'''
p = doc.add_paragraph(code14)
p.style = 'Quote'

doc.add_paragraph()

# --- POST /v1/subscriptions/cancel ---
doc.add_heading('POST /v1/subscriptions/cancel', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Cancel active subscription')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code15 = '''{
    "reason": "No longer need the subscription"
}'''
p = doc.add_paragraph(code15)
p.style = 'Quote'

doc.add_paragraph()

# --- POST /v1/subscriptions/pause ---
doc.add_heading('POST /v1/subscriptions/pause', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Pause active subscription')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- POST /v1/subscriptions/resume ---
doc.add_heading('POST /v1/subscriptions/resume', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Resume paused subscription')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 8. NOTIFICATIONS
# ============================================================
doc.add_heading('8. NOTIFICATIONS', level=1)

# --- POST /v1/notifications/devices ---
doc.add_heading('POST /v1/notifications/devices', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Register device token for push notifications')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code16 = '''{
    "token": "device_token_here",
    "platform": "android"
}'''
p = doc.add_paragraph(code16)
p.style = 'Quote'

doc.add_paragraph()

# --- DELETE /v1/notifications/devices/{token} ---
doc.add_heading('DELETE /v1/notifications/devices/{token}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Remove device token')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()

# --- POST /v1/notifications/test ---
doc.add_heading('POST /v1/notifications/test', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Send test push notification')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-User-ID: {{userId}}')

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 9. CAFES
# ============================================================
doc.add_heading('9. CAFES', level=1)

# --- GET /v1/cafes ---
doc.add_heading('GET /v1/cafes', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get list of cafes')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('companyId: UUID (required)')

doc.add_paragraph()

# --- GET /v1/cafes/{id} ---
doc.add_heading('GET /v1/cafes/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get cafe by ID')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 10. ADMIN - PRODUCTS
# ============================================================
doc.add_heading('10. ADMIN - PRODUCTS', level=1)
p = doc.add_paragraph()
p.add_run('Headers: ').bold = True
p.add_run('X-Company-ID: {{companyId}}')

# --- GET /v1/admin/products ---
doc.add_heading('GET /v1/admin/products', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get products list')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-Company-ID: {{companyId}}')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('categoryId: UUID (optional)')
doc.add_paragraph('status: string (optional)')
doc.add_paragraph('page: int (default: 0)')
doc.add_paragraph('size: int (default: 20)')

doc.add_paragraph()

# --- GET /v1/admin/products/{id} ---
doc.add_heading('GET /v1/admin/products/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get product by ID')

doc.add_paragraph()

# --- POST /v1/admin/products ---
doc.add_heading('POST /v1/admin/products', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Create new product')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-Company-ID: {{companyId}}')
doc.add_paragraph('Content-Type: application/json')
p = doc.add_paragraph()
p.add_run('Body (JSON):').bold = True

code17 = '''{
    "categoryId": "REPLACE_WITH_CATEGORY_ID",
    "name": "Ca Phe Sua",
    "description": "Ca phe sua da dam da",
    "shortDescription": "Ca phe sua",
    "price": 35000,
    "originalPrice": 40000,
    "imageUrl": "https://example.com/coffee.jpg",
    "stock": 100
}'''
p = doc.add_paragraph(code17)
p.style = 'Quote'

doc.add_paragraph()

# --- PUT /v1/admin/products/{id} ---
doc.add_heading('PUT /v1/admin/products/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Update product')
doc.add_paragraph('Content-Type: application/json')

doc.add_paragraph()

# --- DELETE /v1/admin/products/{id} ---
doc.add_heading('DELETE /v1/admin/products/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Delete product')

doc.add_paragraph()

# --- GET /v1/admin/products/categories ---
doc.add_heading('GET /v1/admin/products/categories', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get product categories')

doc.add_paragraph()

# --- POST /v1/admin/products/categories ---
doc.add_heading('POST /v1/admin/products/categories', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Create new category')
doc.add_paragraph('Content-Type: application/json')

code18 = '''{
    "name": "Ca Phe",
    "icon": "coffee",
    "color": "#8B4513",
    "sortOrder": 1
}'''
p = doc.add_paragraph(code18)
p.style = 'Quote'

doc.add_paragraph()

# --- PUT /v1/admin/products/categories/{id} ---
doc.add_heading('PUT /v1/admin/products/categories/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Update category')

doc.add_paragraph()

# --- DELETE /v1/admin/products/categories/{id} ---
doc.add_heading('DELETE /v1/admin/products/categories/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Delete category')

doc.add_paragraph()

# --- GET /v1/admin/products/promotions ---
doc.add_heading('GET /v1/admin/products/promotions', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get promotions')

doc.add_paragraph()

# --- POST /v1/admin/products/promotions ---
doc.add_heading('POST /v1/admin/products/promotions', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Create new promotion')
doc.add_paragraph('Content-Type: application/json')

code19 = '''{
    "code": "SUMMER2024",
    "name": "Summer Sale",
    "description": "Summer promotion 20% off",
    "discountType": "percentage",
    "discountValue": 20,
    "minOrderAmount": 50000,
    "maxDiscountAmount": 20000,
    "startDate": "2024-06-01T00:00:00Z",
    "endDate": "2024-08-31T23:59:59Z",
    "usageLimit": 100
}'''
p = doc.add_paragraph(code19)
p.style = 'Quote'

doc.add_paragraph()

# --- GET /v1/admin/products/rewards ---
doc.add_heading('GET /v1/admin/products/rewards', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get loyalty rewards')

doc.add_paragraph()

# --- POST /v1/admin/products/rewards ---
doc.add_heading('POST /v1/admin/products/rewards', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Create new loyalty reward')
doc.add_paragraph('Content-Type: application/json')

code20 = '''{
    "code": "FREESHIP",
    "name": "Free Shipping",
    "description": "Get free shipping on your order",
    "rewardType": "free_delivery",
    "discountType": "free_delivery",
    "discountValue": 0,
    "pointsCost": 100
}'''
p = doc.add_paragraph(code20)
p.style = 'Quote'

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 11. ADMIN - ORDERS
# ============================================================
doc.add_heading('11. ADMIN - ORDERS', level=1)

# --- GET /v1/admin/orders ---
doc.add_heading('GET /v1/admin/orders', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get orders for admin')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-Company-ID: {{companyId}}')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('cafeId: UUID (optional)')
doc.add_paragraph('status: string (optional)')
doc.add_paragraph('page: int (default: 0)')
doc.add_paragraph('size: int (default: 20)')

doc.add_paragraph()

# --- GET /v1/admin/orders/{id} ---
doc.add_heading('GET /v1/admin/orders/{id}', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get order details for admin')

doc.add_paragraph()

# --- PUT /v1/admin/orders/{id}/status ---
doc.add_heading('PUT /v1/admin/orders/{id}/status', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Update order status')
doc.add_paragraph('Content-Type: application/json')

code21 = '''{
    "status": "preparing",
    "staffNote": "Order in progress"
}'''
p = doc.add_paragraph(code21)
p.style = 'Quote'

doc.add_paragraph()

# --- PUT /v1/admin/orders/{id}/assign ---
doc.add_heading('PUT /v1/admin/orders/{id}/assign', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Assign order to staff/cafe')
doc.add_paragraph('Content-Type: application/json')

code22 = '''{
    "staffId": "REPLACE_WITH_STAFF_ID",
    "cafeId": "REPLACE_WITH_CAFE_ID"
}'''
p = doc.add_paragraph(code22)
p.style = 'Quote'

doc.add_paragraph()

# --- PUT /v1/admin/orders/{id}/note ---
doc.add_heading('PUT /v1/admin/orders/{id}/note', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Update order note')
doc.add_paragraph('Content-Type: application/json')

code23 = '''{
    "note": "Customer VIP"
}'''
p = doc.add_paragraph(code23)
p.style = 'Quote'

doc.add_paragraph()

# --- PUT /v1/admin/orders/{id}/items/{itemId}/status ---
doc.add_heading('PUT /v1/admin/orders/{id}/items/{itemId}/status', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Update order item status')
doc.add_paragraph('Content-Type: application/json')

code24 = '''{
    "status": "ready"
}'''
p = doc.add_paragraph(code24)
p.style = 'Quote'

doc.add_paragraph()

# --- POST /v1/admin/orders/{id}/cancel ---
doc.add_heading('POST /v1/admin/orders/{id}/cancel', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Cancel order as admin')
doc.add_paragraph('Content-Type: application/json')

code25 = '''{
    "reason": "Out of stock"
}'''
p = doc.add_paragraph(code25)
p.style = 'Quote'

doc.add_paragraph()

# --- POST /v1/admin/orders/{id}/refund ---
doc.add_heading('POST /v1/admin/orders/{id}/refund', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Refund an order')
doc.add_paragraph('Content-Type: application/json')

code26 = '''{
    "reason": "Customer request"
}'''
p = doc.add_paragraph(code26)
p.style = 'Quote'

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# 12. ADMIN - DASHBOARD
# ============================================================
doc.add_heading('12. ADMIN - DASHBOARD', level=1)

# --- GET /v1/admin/dashboard/overview ---
doc.add_heading('GET /v1/admin/dashboard/overview', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get dashboard overview statistics')
p = doc.add_paragraph()
p.add_run('Headers:').bold = True
doc.add_paragraph('Authorization: Bearer {{accessToken}}')
doc.add_paragraph('X-Company-ID: {{companyId}}')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('cafeId: UUID (optional)')
doc.add_paragraph('period: string (default: "month") - day|week|month|year')

doc.add_paragraph()

# --- GET /v1/admin/dashboard/revenue-chart ---
doc.add_heading('GET /v1/admin/dashboard/revenue-chart', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get revenue chart data')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('cafeId: UUID (optional)')
doc.add_paragraph('period: string (default: "month")')

doc.add_paragraph()

# --- GET /v1/admin/dashboard/top-products ---
doc.add_heading('GET /v1/admin/dashboard/top-products', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get top selling products')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('cafeId: UUID (optional)')
doc.add_paragraph('limit: int (default: 10)')

doc.add_paragraph()

# --- GET /v1/admin/dashboard/quick-stats ---
doc.add_heading('GET /v1/admin/dashboard/quick-stats', level=2)
p = doc.add_paragraph()
p.add_run('Description: ').bold = True
p.add_run('Get quick statistics')
p = doc.add_paragraph()
p.add_run('Query Params:').bold = True
doc.add_paragraph('cafeId: UUID (optional)')

doc.add_paragraph()
doc.add_paragraph('=' * 80)

# ============================================================
# PLACEHOLDERS
# ============================================================
doc.add_heading('PLACEHOLDERS', level=1)

table = doc.add_table(rows=11, cols=2)
table.style = 'Table Grid'

hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Placeholder'
hdr_cells[1].text = 'Description'

data = [
    ('{{accessToken}}', 'JWT access token (auto saved after login)'),
    ('{{refreshToken}}', 'Refresh token (auto saved after login)'),
    ('{{userId}}', 'User ID (auto saved after login)'),
    ('{{companyId}}', 'Company UUID'),
    ('REPLACE_WITH_CAFE_ID', 'Cafe UUID'),
    ('REPLACE_WITH_PRODUCT_ID', 'Product UUID'),
    ('REPLACE_WITH_ORDER_ID', 'Order UUID'),
    ('REPLACE_WITH_CATEGORY_ID', 'Category UUID'),
    ('REPLACE_WITH_REWARD_ID', 'Reward UUID'),
    ('REPLACE_WITH_PLAN_ID', 'Subscription plan UUID'),
]

for i, (placeholder, desc) in enumerate(data, 1):
    row = table.rows[i].cells
    row[0].text = placeholder
    row[1].text = desc

doc.add_paragraph()

# ============================================================
# RESPONSE FORMAT
# ============================================================
doc.add_heading('RESPONSE FORMAT', level=1)

code_response = '''{
    "success": true,
    "data": { ... },
    "message": "Success",
    "timestamp": "2024-01-01T12:00:00Z"
}'''
p = doc.add_paragraph(code_response)
p.style = 'Quote'

# ============================================================
# SAVE
# ============================================================
doc.save('D:/project code/AI Cafe Platform/BackEnd/java/Postman Collection/AI_Cafe_API_Documentation.docx')
print('Created: AI_Cafe_API_Documentation.docx')
