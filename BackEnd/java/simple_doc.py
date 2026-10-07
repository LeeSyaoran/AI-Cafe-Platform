from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# TITLE
title = doc.add_heading('AI CAFE PLATFORM - API DOCUMENTATION', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Base URL: https://api-staging.aicafe.vn/api')
doc.add_paragraph('Production: https://api.aicafe.vn/api')
doc.add_paragraph('Local: http://localhost:3000/api')
doc.add_paragraph('=' * 80)

# 1. AUTHENTICATION
doc.add_heading('1. AUTHENTICATION', level=1)

doc.add_heading('POST /v1/auth/send-otp', level=2)
doc.add_paragraph('Description: Send OTP to phone or email')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"email": "test@example.com", "phone": "0909123456", "type": "register"}')
p.style = 'Quote'

doc.add_heading('POST /v1/auth/verify-otp', level=2)
doc.add_paragraph('Description: Verify OTP and get tokens')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"email": "test@example.com", "phone": "0909123456", "otp": "123456"}')
p.style = 'Quote'

doc.add_heading('POST /v1/auth/register', level=2)
doc.add_paragraph('Description: Register new user account')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"email": "test@example.com", "password": "password123", "phone": "0909123456"}')
p.style = 'Quote'

doc.add_heading('POST /v1/auth/login', level=2)
doc.add_paragraph('Description: Login with email and password')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"email": "test@example.com", "password": "password123"}')
p.style = 'Quote'

doc.add_heading('POST /v1/auth/refresh', level=2)
doc.add_paragraph('Description: Refresh access token')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"refreshToken": "{{refreshToken}}"}')
p.style = 'Quote'

doc.add_heading('POST /v1/auth/logout', level=2)
doc.add_paragraph('Description: Logout and invalidate token')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"refreshToken": "{{refreshToken}}"}')
p.style = 'Quote'

doc.add_paragraph('=' * 80)

# 2. USER PROFILE
doc.add_heading('2. USER PROFILE', level=1)

doc.add_heading('GET /v1/me', level=2)
doc.add_paragraph('Description: Get current user profile')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('PUT /v1/me', level=2)
doc.add_paragraph('Description: Update current user profile')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"fullName": "Nguyen Van Test", "displayName": "TestUser", "dateOfBirth": "1990-01-15", "gender": "male"}')
p.style = 'Quote'

doc.add_heading('GET /v1/users/{id}', level=2)
doc.add_paragraph('Description: Get user by ID')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_paragraph('=' * 80)

# 3. CART
doc.add_heading('3. CART', level=1)

doc.add_heading('GET /v1/cart', level=2)
doc.add_paragraph('Description: Get user cart for a cafe')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Query: cafeId=UUID')

doc.add_heading('POST /v1/cart/items', level=2)
doc.add_paragraph('Description: Add item to cart')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"cafeId": "REPLACE_WITH_CAFE_ID", "productId": "REPLACE_WITH_PRODUCT_ID", "quantity": 1, "options": {}, "modifiers": [], "notes": ""}')
p.style = 'Quote'

doc.add_heading('PUT /v1/cart/items/{id}', level=2)
doc.add_paragraph('Description: Update cart item')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"quantity": 2, "notes": "Extra ice"}')
p.style = 'Quote'

doc.add_heading('DELETE /v1/cart/items/{id}', level=2)
doc.add_paragraph('Description: Remove item from cart')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('DELETE /v1/cart', level=2)
doc.add_paragraph('Description: Clear all items from cart')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Query: cafeId=UUID')

doc.add_paragraph('=' * 80)

# 4. ORDERS
doc.add_heading('4. ORDERS', level=1)

doc.add_heading('POST /v1/orders', level=2)
doc.add_paragraph('Description: Create new order from cart')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"companyId": "REPLACE_WITH_COMPANY_ID", "cafeId": "REPLACE_WITH_CAFE_ID", "orderType": "pickup", "customerNote": "", "promoCode": ""}')
p.style = 'Quote'

doc.add_heading('GET /v1/orders', level=2)
doc.add_paragraph('Description: Get user order history')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Query: page=0, size=20')

doc.add_heading('GET /v1/orders/{id}', level=2)
doc.add_paragraph('Description: Get order by ID')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_heading('GET /v1/orders/number/{orderNumber}', level=2)
doc.add_paragraph('Description: Get order by order number')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_heading('PUT /v1/orders/{id}/confirm', level=2)
doc.add_paragraph('Description: Staff confirms order')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_heading('PUT /v1/orders/{id}/preparing', level=2)
doc.add_paragraph('Description: Staff starts preparing order')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_heading('PUT /v1/orders/{id}/ready', level=2)
doc.add_paragraph('Description: Staff marks order as ready')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_heading('PUT /v1/orders/{id}/complete', level=2)
doc.add_paragraph('Description: Staff completes order')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_heading('POST /v1/orders/{id}/cancel', level=2)
doc.add_paragraph('Description: Cancel order')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"reason": "Customer requested cancellation"}')
p.style = 'Quote'

doc.add_paragraph('=' * 80)

# 5. PAYMENTS
doc.add_heading('5. PAYMENTS', level=1)

doc.add_heading('GET /v1/payments/methods', level=2)
doc.add_paragraph('Description: Get available payment methods')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('POST /v1/payments', level=2)
doc.add_paragraph('Description: Create payment for order')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"orderId": "REPLACE_WITH_ORDER_ID", "companyId": "REPLACE_WITH_COMPANY_ID", "amount": 50000, "paymentMethod": "CREDIT", "returnUrl": "https://aicafe.vn/payment/return"}')
p.style = 'Quote'

doc.add_heading('GET /v1/payments/{id}', level=2)
doc.add_paragraph('Description: Get payment by ID')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('GET /v1/payments/order/{orderId}', level=2)
doc.add_paragraph('Description: Get payment by order ID')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('POST /v1/payments/{id}/cancel', level=2)
doc.add_paragraph('Description: Cancel pending payment')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_paragraph('=' * 80)

# 6. LOYALTY & REWARDS
doc.add_heading('6. LOYALTY & REWARDS', level=1)

doc.add_heading('GET /v1/loyalty/account', level=2)
doc.add_paragraph('Description: Get user loyalty account info')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('GET /v1/loyalty/rewards', level=2)
doc.add_paragraph('Description: Get available rewards for redemption')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('POST /v1/loyalty/redeem', level=2)
doc.add_paragraph('Description: Redeem loyalty points for reward')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"rewardId": "REPLACE_WITH_REWARD_ID", "orderId": "REPLACE_WITH_ORDER_ID"}')
p.style = 'Quote'

doc.add_heading('GET /v1/loyalty/history', level=2)
doc.add_paragraph('Description: Get loyalty transaction history')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Query: limit=50')

doc.add_heading('GET /v1/loyalty/redemptions', level=2)
doc.add_paragraph('Description: Get reward redemption history')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_paragraph('=' * 80)

# 7. SUBSCRIPTIONS
doc.add_heading('7. SUBSCRIPTIONS', level=1)

doc.add_heading('GET /v1/subscriptions/plans', level=2)
doc.add_paragraph('Description: Get available membership plans')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_heading('GET /v1/subscriptions/plans/{id}', level=2)
doc.add_paragraph('Description: Get subscription plan details')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_heading('GET /v1/subscriptions/me', level=2)
doc.add_paragraph('Description: Get user current subscription')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('POST /v1/subscriptions', level=2)
doc.add_paragraph('Description: Subscribe to membership plan')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"planId": "REPLACE_WITH_PLAN_ID", "billingCycle": "monthly", "paymentMethod": "CREDIT"}')
p.style = 'Quote'

doc.add_heading('POST /v1/subscriptions/cancel', level=2)
doc.add_paragraph('Description: Cancel active subscription')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"reason": "No longer need the subscription"}')
p.style = 'Quote'

doc.add_heading('POST /v1/subscriptions/pause', level=2)
doc.add_paragraph('Description: Pause active subscription')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('POST /v1/subscriptions/resume', level=2)
doc.add_paragraph('Description: Resume paused subscription')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_paragraph('=' * 80)

# 8. NOTIFICATIONS
doc.add_heading('8. NOTIFICATIONS', level=1)

doc.add_heading('POST /v1/notifications/devices', level=2)
doc.add_paragraph('Description: Register device token for push notifications')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"token": "device_token_here", "platform": "android"}')
p.style = 'Quote'

doc.add_heading('DELETE /v1/notifications/devices/{token}', level=2)
doc.add_paragraph('Description: Remove device token')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_heading('POST /v1/notifications/test', level=2)
doc.add_paragraph('Description: Send test push notification')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-User-ID: {{userId}}')

doc.add_paragraph('=' * 80)

# 9. CAFES
doc.add_heading('9. CAFES', level=1)

doc.add_heading('GET /v1/cafes', level=2)
doc.add_paragraph('Description: Get list of cafes')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')
doc.add_paragraph('Query: companyId=UUID')

doc.add_heading('GET /v1/cafes/{id}', level=2)
doc.add_paragraph('Description: Get cafe by ID')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}')

doc.add_paragraph('=' * 80)

# 10. ADMIN - PRODUCTS
doc.add_heading('10. ADMIN - PRODUCTS', level=1)
doc.add_paragraph('Headers: X-Company-ID: {{companyId}}')

doc.add_heading('GET /v1/admin/products', level=2)
doc.add_paragraph('Description: Get products list')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-Company-ID: {{companyId}}')
doc.add_paragraph('Query: categoryId, status, page=0, size=20')

doc.add_heading('GET /v1/admin/products/{id}', level=2)
doc.add_paragraph('Description: Get product by ID')

doc.add_heading('POST /v1/admin/products', level=2)
doc.add_paragraph('Description: Create new product')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-Company-ID: {{companyId}}')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"categoryId": "REPLACE_WITH_CATEGORY_ID", "name": "Ca Phe Sua", "description": "Ca phe sua da dam da", "price": 35000, "originalPrice": 40000, "imageUrl": "https://example.com/coffee.jpg", "stock": 100}')
p.style = 'Quote'

doc.add_heading('PUT /v1/admin/products/{id}', level=2)
doc.add_paragraph('Description: Update product')

doc.add_heading('DELETE /v1/admin/products/{id}', level=2)
doc.add_paragraph('Description: Delete product')

doc.add_heading('GET /v1/admin/products/categories', level=2)
doc.add_paragraph('Description: Get product categories')

doc.add_heading('POST /v1/admin/products/categories', level=2)
doc.add_paragraph('Description: Create new category')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"name": "Ca Phe", "icon": "coffee", "color": "#8B4513", "sortOrder": 1}')
p.style = 'Quote'

doc.add_heading('PUT /v1/admin/products/categories/{id}', level=2)
doc.add_paragraph('Description: Update category')

doc.add_heading('DELETE /v1/admin/products/categories/{id}', level=2)
doc.add_paragraph('Description: Delete category')

doc.add_heading('GET /v1/admin/products/promotions', level=2)
doc.add_paragraph('Description: Get promotions')

doc.add_heading('POST /v1/admin/products/promotions', level=2)
doc.add_paragraph('Description: Create new promotion')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"code": "SUMMER2024", "name": "Summer Sale", "discountType": "percentage", "discountValue": 20, "minOrderAmount": 50000, "maxDiscountAmount": 20000, "startDate": "2024-06-01T00:00:00Z", "endDate": "2024-08-31T23:59:59Z", "usageLimit": 100}')
p.style = 'Quote'

doc.add_heading('GET /v1/admin/products/rewards', level=2)
doc.add_paragraph('Description: Get loyalty rewards')

doc.add_heading('POST /v1/admin/products/rewards', level=2)
doc.add_paragraph('Description: Create new loyalty reward')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"code": "FREESHIP", "name": "Free Shipping", "rewardType": "free_delivery", "pointsCost": 100}')
p.style = 'Quote'

doc.add_paragraph('=' * 80)

# 11. ADMIN - ORDERS
doc.add_heading('11. ADMIN - ORDERS', level=1)

doc.add_heading('GET /v1/admin/orders', level=2)
doc.add_paragraph('Description: Get orders for admin')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-Company-ID: {{companyId}}')
doc.add_paragraph('Query: cafeId, status, page=0, size=20')

doc.add_heading('GET /v1/admin/orders/{id}', level=2)
doc.add_paragraph('Description: Get order details for admin')

doc.add_heading('PUT /v1/admin/orders/{id}/status', level=2)
doc.add_paragraph('Description: Update order status')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"status": "preparing", "staffNote": "Order in progress"}')
p.style = 'Quote'

doc.add_heading('PUT /v1/admin/orders/{id}/assign', level=2)
doc.add_paragraph('Description: Assign order to staff/cafe')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"staffId": "REPLACE_WITH_STAFF_ID", "cafeId": "REPLACE_WITH_CAFE_ID"}')
p.style = 'Quote'

doc.add_heading('PUT /v1/admin/orders/{id}/note', level=2)
doc.add_paragraph('Description: Update order note')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"note": "Customer VIP"}')
p.style = 'Quote'

doc.add_heading('PUT /v1/admin/orders/{id}/items/{itemId}/status', level=2)
doc.add_paragraph('Description: Update order item status')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"status": "ready"}')
p.style = 'Quote'

doc.add_heading('POST /v1/admin/orders/{id}/cancel', level=2)
doc.add_paragraph('Description: Cancel order as admin')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"reason": "Out of stock"}')
p.style = 'Quote'

doc.add_heading('POST /v1/admin/orders/{id}/refund', level=2)
doc.add_paragraph('Description: Refund an order')
doc.add_paragraph('Body:')
p = doc.add_paragraph('{"reason": "Customer request"}')
p.style = 'Quote'

doc.add_paragraph('=' * 80)

# 12. ADMIN - DASHBOARD
doc.add_heading('12. ADMIN - DASHBOARD', level=1)

doc.add_heading('GET /v1/admin/dashboard/overview', level=2)
doc.add_paragraph('Description: Get dashboard overview statistics')
doc.add_paragraph('Headers: Authorization: Bearer {{accessToken}}, X-Company-ID: {{companyId}}')
doc.add_paragraph('Query: cafeId, period=month (day|week|month|year)')

doc.add_heading('GET /v1/admin/dashboard/revenue-chart', level=2)
doc.add_paragraph('Description: Get revenue chart data')
doc.add_paragraph('Query: cafeId, period=month')

doc.add_heading('GET /v1/admin/dashboard/top-products', level=2)
doc.add_paragraph('Description: Get top selling products')
doc.add_paragraph('Query: cafeId, limit=10')

doc.add_heading('GET /v1/admin/dashboard/quick-stats', level=2)
doc.add_paragraph('Description: Get quick statistics')
doc.add_paragraph('Query: cafeId')

doc.add_paragraph('=' * 80)

# PLACEHOLDERS
doc.add_heading('PLACEHOLDERS', level=1)

table = doc.add_table(rows=11, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Placeholder'
hdr_cells[1].text = 'Description'

data = [
    ('{{accessToken}}', 'JWT access token (auto saved after login)'),
    ('{{refreshToken}}', 'Refresh token (auto saved after login)'),
    ('{{userId}}', 'User ID (auto saved after login)'),
    ('{{companyId}}', 'Company UUID'),
    ('REPLACE_WITH_CAFE_ID', 'Cafe UUID'),
    ('REPLACE_WITH_PRODUCT_ID', 'Product UUID'),
    ('REPLACE_WITH_ORDER_ID', 'Order UUID'),
    ('REPLACE_WITH_CATEGORY_ID', 'Category UUID'),
    ('REPLACE_WITH_REWARD_ID', 'Reward UUID'),
    ('REPLACE_WITH_PLAN_ID', 'Subscription plan UUID'),
]

for i, (placeholder, desc) in enumerate(data, 1):
    row = table.rows[i].cells
    row[0].text = placeholder
    row[1].text = desc

doc.add_paragraph()

# RESPONSE FORMAT
doc.add_heading('RESPONSE FORMAT', level=1)
p = doc.add_paragraph('{"success": true, "data": {...}, "message": "Success", "timestamp": "2024-01-01T12:00:00Z"}')
p.style = 'Quote'

# SAVE
doc.save('D:/project code/AI Cafe Platform/BackEnd/java/Postman Collection/AI_Cafe_API_Documentation.docx')
print('Created: AI_Cafe_API_Documentation.docx')
